"""
Chuyển file Excel quản lý thu chi (mẫu "BẢNG THEO DÕI THU CHI") thành file SQL
để nhập vào Supabase: giao dịch (sheet "Thu Chi"), ngân sách (sheet "Lập ngân
sách") và danh mục khoản mục (sheet "Phân loại thu chi").

Cách dùng:
    pip install openpyxl
    python scripts/excel-to-sql.py "doc/T10-T12-2026 (1).xlsx"
    python scripts/excel-to-sql.py "file.xlsx" --out supabase/seed.local.sql --year 2026

Sau đó mở Supabase → SQL Editor, dán nội dung file .sql vừa tạo và bấm Run
(chạy SAU khi đã chạy supabase/schema.sql và đã tạo tài khoản đăng nhập).
"""

import argparse
import datetime as dt
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Thiếu thư viện openpyxl. Cài bằng: pip install openpyxl")

INCOME_CATEGORY = {
    "Thu nhập chính": "main",
    "Thu nhập phụ": "side",
    "Thu nhập bất thường": "irregular",
}
EXPENSE_CATEGORY = {
    "Thiết yếu": "essential",
    "Không thiết yếu": "non_essential",
    "Phát sinh": "incidental",
    "Tiết kiệm/Đầu tư": "saving",
}
# Phân loại gợi ý cho các khoản mục trong danh sách nhưng chưa từng được dùng.
GUESSED_CATEGORY = {
    "Tiết kiệm": "saving",
    "Học": "essential",
    "Sức khỏe": "essential",
    "Khám bệnh": "essential",
    "Sách/Tài liệu học": "essential",
    "Trà sữa/Cà phê": "non_essential",
    "Sửa xe máy": "incidental",
    "Quà sinh nhật bạn": "incidental",
    "Quà Tết/Biếu gia đình": "incidental",
    "Quà Tết/Lì xì": "incidental",
}
# Sheet "Lập ngân sách": dòng → mã phân loại (cột C..N = tháng 1..12)
BUDGET_ROWS = {5: "main", 6: "side", 7: "irregular", 26: "essential", 27: "non_essential", 28: "incidental", 47: "saving"}


def sql_text(value):
    if value is None or value == "":
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def clean(value):
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def as_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def read_transactions(wb):
    ws = wb["Thu Chi"]
    rows, problems = [], []
    for r in range(5, ws.max_row + 1):
        # Thu nhập: A..E, Chi tiêu: G..K
        for kind, first_col, mapping in (("income", 1, INCOME_CATEGORY), ("expense", 7, EXPENSE_CATEGORY)):
            name, date, category, amount, note = (ws.cell(r, first_col + i).value for i in range(5))
            if name is None and amount is None:
                continue
            day = as_date(date)
            code = mapping.get(clean(category))
            try:
                value = int(round(float(amount)))
            except (TypeError, ValueError):
                value = 0
            if not clean(name) or day is None or code is None or value <= 0:
                problems.append(f"  dòng {r} ({'thu' if kind == 'income' else 'chi'}): {name!r}, {date!r}, {category!r}, {amount!r}")
                continue
            rows.append({
                "kind": kind,
                "item": clean(name),
                "category": code,
                "amount": value,
                "date": day,
                "note": clean(note) or None,
            })
    return rows, problems


def read_budget(wb):
    ws = wb["Lập ngân sách"]
    budget = []
    for row, code in BUDGET_ROWS.items():
        for month in range(1, 13):
            value = ws.cell(row, 2 + month).value
            if isinstance(value, (int, float)) and value > 0:
                budget.append((month, code, int(round(value))))
    return budget


def detect_year(wb, transactions):
    ws = wb["Phân loại thu chi"]
    for row in ws.iter_rows(min_row=1, max_row=40):
        for cell in row:
            if isinstance(cell.value, str) and "CẢ NĂM" in cell.value.upper():
                match = re.search(r"(20\d{2})", cell.value)
                if match:
                    return int(match.group(1))
    years = Counter(t["date"].year for t in transactions)
    return years.most_common(1)[0][0] if years else dt.date.today().year


def read_items(wb, transactions):
    listed = []
    ws = wb["Phân loại thu chi"]
    for r in range(20, 60):
        name = clean(ws.cell(r, 24).value)  # cột X
        if not name or name.upper().startswith("TỔNG"):
            break
        listed.append(name)

    used = defaultdict(list)
    for t in sorted(transactions, key=lambda t: t["date"]):
        used[(t["kind"], t["item"])].append(t["category"])

    items = {}
    for name in listed:
        items[("expense", name)] = None
    for key in used:
        items.setdefault(key, None)
    for key in items:
        categories = used.get(key, [])
        if categories:
            # Chỉ đặt mặc định khi luôn dùng cùng một phân loại; nếu không, app tự lấy phân loại dùng gần nhất.
            items[key] = categories[0] if len(set(categories)) == 1 else None
        else:
            items[key] = GUESSED_CATEGORY.get(key[1])
    return items


def build_sql(source_name, year, transactions, budget, items):
    out = []
    w = out.append
    w(f"-- Dữ liệu nhập từ file Excel: {source_name}")
    w(f"-- {len(transactions)} giao dịch, {len(budget)} ô ngân sách (năm {year}), {len(items)} khoản mục.")
    w("-- Chạy SAU supabase/schema.sql và SAU khi đã tạo tài khoản (Authentication → Users → Add user).")
    w("-- Chạy lại nhiều lần không bị nhân đôi giao dịch.")
    w("do $$")
    w("declare")
    w("  v_user uuid;")
    w("begin")
    w("  if (select count(*) from auth.users) <> 1 then")
    w("    raise exception 'Cần đúng 1 tài khoản trong Authentication → Users (hiện có %). Nếu có nhiều tài khoản, sửa câu lệnh chọn v_user theo email.', (select count(*) from auth.users);")
    w("  end if;")
    w("  select id into v_user from auth.users limit 1;")
    w("")
    w("  if exists (select 1 from public.transactions where user_id = v_user) then")
    w("    raise notice 'Đã có giao dịch trong database — bỏ qua phần giao dịch để tránh trùng lặp.';")
    if transactions:
        w("  else")
        w("    insert into public.transactions (user_id, kind, item, category, amount, occurred_on, note) values")
        lines = [
            f"      (v_user, '{t['kind']}', {sql_text(t['item'])}, '{t['category']}', {t['amount']}, '{t['date'].isoformat()}', {sql_text(t['note'])})"
            for t in transactions
        ]
        w(",\n".join(lines) + ";")
    w("  end if;")
    if budget:
        w("")
        w("  insert into public.budgets (user_id, year, month, category, amount) values")
        w(",\n".join(f"    (v_user, {year}, {m}, '{c}', {a})" for m, c, a in budget))
        w("  on conflict (user_id, year, month, category) do update set amount = excluded.amount;")
    if items:
        w("")
        w("  insert into public.items (user_id, kind, name, default_category) values")
        w(",\n".join(
            f"    (v_user, '{kind}', {sql_text(name)}, {sql_text(category)})"
            for (kind, name), category in items.items()
        ))
        w("  on conflict (user_id, kind, name) do nothing;")
    w("end $$;")
    return "\n".join(out) + "\n"


def main():
    parser = argparse.ArgumentParser(description="Excel thu chi → SQL cho Supabase")
    parser.add_argument("xlsx", help="Đường dẫn file Excel")
    parser.add_argument("--out", default="supabase/seed.local.sql", help="File SQL đầu ra")
    parser.add_argument("--year", type=int, help="Năm của sheet Lập ngân sách (mặc định: tự nhận diện)")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    transactions, problems = read_transactions(wb)
    year = args.year or detect_year(wb, transactions)
    budget = read_budget(wb)
    items = read_items(wb, transactions)

    sql = build_sql(Path(args.xlsx).name, year, transactions, budget, items)
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    Path(args.out).write_text(sql, encoding="utf-8")

    print(f"Đã ghi {args.out}: {len(transactions)} giao dịch, {len(budget)} ô ngân sách năm {year}, {len(items)} khoản mục.")
    if problems:
        print("Bỏ qua các dòng thiếu/không hợp lệ:")
        print("\n".join(problems))


if __name__ == "__main__":
    main()
